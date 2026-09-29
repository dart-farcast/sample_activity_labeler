"""
Sample Activity Labeler
========================
Converts the multimodal-assay analysis notebook into an interactive
Streamlit app for labeling sample-level biological activity.

Pipeline:
1. Load 4 assays (Histo, Cytokine, Flow, NanoString) from the multimodal
   Excel workbook, plus a Gene Signature (GES) list CSV.
2. Compute NanoString Gene Expression Signature (GES) scores.
3. Compute log2 fold-change (user-selected Treatment vs. Control arm) for
   every assay.
4. Select "responsive" features (|log2FC| >= cutoff in >= N samples).
5. Classify each selected feature's direction (Increased / Decreased / No
   meaningful change) per sample.
6. Score each biological signature (Histo / Cytokine / Flow) as
   Active / No clear activity / Insufficient data, and derive NanoString
   activity directly from the GES direction table.
7. Collapse each assay's activity into one label per sample.
8. Final output: a binary table where 1 = that activity signature is
   present (active) for that sample, 0 = not present.

Run with:
    streamlit run rxb_activity_app.py
"""

import base64
import io
import re
from datetime import datetime
from pathlib import Path

import numpy as np
import pandas as pd
import streamlit as st

# ============================================================
# FARCAST BRAND ASSETS & THEME SETUP
# ============================================================
LOGO_PATH = Path(__file__).parent / "farcast biodynamic logo.jpeg"
LOGO_B64 = ""
if LOGO_PATH.exists():
    try:
        LOGO_B64 = base64.b64encode(LOGO_PATH.read_bytes()).decode("utf-8")
    except Exception:
        LOGO_B64 = ""

st.set_page_config(
    page_title="Farcast Biosciences | Sample Activity Labeler",
    page_icon="🧬",
    layout="wide",
    initial_sidebar_state="expanded",
)

FARCAST_CSS = """
<style>
    @import url('https://fonts.googleapis.com/css2?family=Roboto+Condensed:wght@400;600;700&family=Roboto:ital,wght@0,300;0,400;0,500;0,700;1,400&display=swap');

    :root {
        --farcast-navy: #1E2859;
        --farcast-navy-dark: #131A3D;
        --farcast-navy-light: #2A3875;
        --farcast-purple: #B14FC4;
        --farcast-purple-dark: #8E2FA0;
        --farcast-purple-light: #F7EBFB;
        --farcast-purple-subtle: #FAF3FC;
        --farcast-border: #E2E8F0;
        --farcast-bg-card: #FFFFFF;
        --farcast-bg-subtle: #F8F9FD;
        --farcast-text-main: #1E2859;
        --farcast-text-muted: #5A6478;
    }

    /* Global Typography */
    html, body, [class*="css"] {
        font-family: 'Roboto', -apple-system, BlinkMacSystemFont, 'Segoe UI', Arial, sans-serif;
        color: var(--farcast-text-main);
    }

    /* Classic app canvas */
    .block-container,
    [data-testid="stMainBlockContainer"] {
        padding-top: 1rem !important;
        padding-bottom: 2.5rem !important;
        padding-left: 1.8rem !important;
        padding-right: 1.8rem !important;
        max-width: 1040px !important;
        margin-left: auto !important;
        margin-right: auto !important;
    }

    /* Cleaner app chrome */
    #MainMenu, footer { visibility: hidden; }
    header[data-testid="stHeader"] { background: transparent; }

    /* Typography */
    h1, h2, h3, h4 {
        font-family: 'Roboto Condensed', 'Roboto', sans-serif !important;
        letter-spacing: -0.01em;
        color: var(--farcast-navy) !important;
    }
    h1 {
        font-weight: 700 !important;
    }
    h2 {
        font-weight: 700 !important;
        font-size: 1.35rem !important;
        border-bottom: 2px solid var(--farcast-purple-light);
        padding-bottom: 0.35rem;
        margin-top: 1.2rem !important;
    }
    h3 {
        font-weight: 600 !important;
        font-size: 1.1rem !important;
        color: var(--farcast-navy-light) !important;
    }
    .app-subtitle {
        color: var(--farcast-text-muted);
        font-size: 0.92rem;
        line-height: 1.45;
        margin-top: -0.2rem;
        margin-bottom: 1.2rem;
    }

    /* Header banner */
    .farcast-header-card {
        background-color: #1E2859;
        border-radius: 10px;
        padding: 1.3rem 1.7rem;
        color: #FFFFFF;
        box-shadow: 0 4px 16px rgba(30, 40, 89, 0.12);
        margin-bottom: 1.2rem;
        position: relative;
        border-left: 8px solid #B14FC4;
        border-bottom: 2px solid #B14FC4;
    }
    .farcast-header-brand {
        display: inline-flex;
        align-items: center;
        gap: 0.45rem;
        background-color: #B14FC4;
        padding: 0.25rem 0.75rem;
        border-radius: 4px;
        font-size: 0.75rem;
        font-weight: 700;
        letter-spacing: 0.08em;
        text-transform: uppercase;
        color: #FFFFFF;
        margin-bottom: 0.6rem;
    }
    .farcast-header-title {
        color: #FFFFFF !important;
        font-size: 1.55rem !important;
        font-weight: 700 !important;
        margin: 0 0 0.3rem 0 !important;
        letter-spacing: -0.01em;
        font-family: 'Roboto Condensed', sans-serif !important;
    }
    .farcast-header-desc {
        color: #E2E6F5 !important;
        margin: 0 !important;
        font-size: 0.85rem !important;
        font-weight: 400;
        line-height: 1.4;
    }
    .farcast-badges {
        display: flex;
        flex-wrap: wrap;
        gap: 0.5rem;
        margin-top: 0.85rem;
    }
    .farcast-badge-chip {
        background-color: rgba(255, 255, 255, 0.1);
        color: #FFFFFF;
        font-size: 0.73rem;
        font-weight: 500;
        padding: 0.22rem 0.65rem;
        border-radius: 4px;
        border: 1px solid #B14FC4;
    }

    /* Sidebar Styling */
    section[data-testid="stSidebar"] {
        background-color: var(--farcast-bg-subtle) !important;
        border-right: 1px solid var(--farcast-border) !important;
    }
    section[data-testid="stSidebar"][aria-expanded="true"] {
        min-width: 270px !important;
        max-width: 270px !important;
    }
    section[data-testid="stSidebar"] .block-container,
    section[data-testid="stSidebar"] [data-testid="stSidebarUserContent"] {
        padding-left: 1rem !important;
        padding-right: 1rem !important;
        max-width: 100% !important;
    }
    .farcast-sidebar-logo-container {
        background: #FFFFFF;
        border-radius: 8px;
        padding: 0.85rem 0.9rem;
        margin-bottom: 1.1rem;
        border: 1px solid var(--farcast-border);
        border-top: 3px solid #1E2859;
        box-shadow: 0 2px 6px rgba(30, 40, 89, 0.04);
        text-align: center;
    }
    .farcast-sidebar-logo {
        max-width: 100%;
        height: auto;
        display: block;
        margin: 0 auto;
    }
    section[data-testid="stSidebar"] h2 {
        font-family: 'Roboto Condensed', sans-serif !important;
        font-size: 0.9rem !important;
        font-weight: 700 !important;
        text-transform: uppercase;
        letter-spacing: 0.05em;
        color: var(--farcast-navy) !important;
        margin-top: 1.1rem !important;
        border-bottom: none !important;
        border-left: 4px solid var(--farcast-purple) !important;
        padding-left: 0.5rem !important;
        padding-bottom: 0 !important;
    }

    /* Buttons */
    .stButton > button, .stDownloadButton > button {
        border-radius: 6px !important;
        font-weight: 600 !important;
        font-family: 'Roboto', sans-serif !important;
        transition: all 0.18s ease-in-out !important;
    }
    .stButton > button[kind="primary"] {
        background-color: #1E2859 !important;
        color: #FFFFFF !important;
        border: 1.5px solid #1E2859 !important;
        box-shadow: 0 2px 6px rgba(30, 40, 89, 0.15) !important;
    }
    .stButton > button[kind="primary"]:hover {
        background-color: #B14FC4 !important;
        border-color: #B14FC4 !important;
        color: #FFFFFF !important;
        box-shadow: 0 4px 12px rgba(177, 79, 196, 0.3) !important;
        transform: translateY(-1px) !important;
    }
    .stDownloadButton > button {
        background-color: #FFFFFF !important;
        color: var(--farcast-navy) !important;
        border: 1.5px solid var(--farcast-navy) !important;
    }
    .stDownloadButton > button:hover {
        background-color: var(--farcast-purple-light) !important;
        border-color: var(--farcast-purple) !important;
        color: var(--farcast-purple-dark) !important;
        transform: translateY(-1px) !important;
    }
    .stDownloadButton > button[kind="primary"] {
        background-color: #1E2859 !important;
        color: #FFFFFF !important;
        border: 1.5px solid #1E2859 !important;
    }
    .stDownloadButton > button[kind="primary"]:hover {
        background-color: #B14FC4 !important;
        border-color: #B14FC4 !important;
        color: #FFFFFF !important;
    }

    /* Metric cards */
    div[data-testid="stMetric"] {
        background-color: #FFFFFF !important;
        border: 1px solid var(--farcast-border) !important;
        border-top: 3.5px solid var(--farcast-navy) !important;
        border-radius: 10px !important;
        padding: 0.85rem 1.1rem 0.65rem 1.1rem !important;
        box-shadow: 0 2px 8px rgba(30, 40, 89, 0.04) !important;
    }
    div[data-testid="stMetric"] label {
        color: var(--farcast-text-muted) !important;
        font-weight: 500 !important;
        font-size: 0.82rem !important;
        text-transform: uppercase !important;
        letter-spacing: 0.03em !important;
    }
    div[data-testid="stMetric"] div[data-testid="stMetricValue"] {
        color: var(--farcast-navy) !important;
        font-family: 'Roboto Condensed', sans-serif !important;
        font-weight: 700 !important;
        font-size: 1.7rem !important;
    }

    /* Dataframes */
    div[data-testid="stDataFrame"] {
        border: 1px solid var(--farcast-border) !important;
        border-radius: 8px !important;
        overflow: hidden !important;
        box-shadow: 0 1px 4px rgba(30, 40, 89, 0.03) !important;
    }

    /* Tabs */
    .stTabs [data-baseweb="tab-list"] {
        gap: 4px;
        border-bottom: 2px solid var(--farcast-border);
    }
    .stTabs [data-baseweb="tab"] {
        font-family: 'Roboto Condensed', sans-serif !important;
        font-weight: 600 !important;
        color: var(--farcast-text-muted) !important;
        border-radius: 6px 6px 0 0 !important;
        padding: 0.4rem 1rem !important;
    }
    .stTabs [aria-selected="true"] {
        color: var(--farcast-navy) !important;
        border-bottom: 3px solid var(--farcast-purple) !important;
        background-color: var(--farcast-purple-light) !important;
    }

    /* Expanders */
    details {
        border: 1px solid var(--farcast-border) !important;
        border-radius: 8px !important;
        background-color: #FFFFFF !important;
    }
    summary {
        font-weight: 600 !important;
        color: var(--farcast-navy) !important;
    }

    /* Interactive Inputs focus */
    .stTextInput input:focus, .stNumberInput input:focus, .stSelectbox select:focus {
        border-color: var(--farcast-purple) !important;
        box-shadow: 0 0 0 2px rgba(177, 79, 196, 0.2) !important;
    }
</style>
"""
st.markdown(FARCAST_CSS, unsafe_allow_html=True)

# ============================================================
# SIGNATURE DEFINITIONS (from the original notebook)
# ============================================================

HISTO_SIGNATURES = {
    "Tumor cytotoxicity": {
        "Median_Tumor_%": -1,
        "Median_Cas-3": 1,
        "Median_Necrosis": 1,
    },
    "Immune infiltration": {
        "Median_Immune_component_%": 1,
        "Median_Tumor_infiltrated_immune_cells_%": 1,
    },
}

CYTOKINE_SIGNATURES = {
    "IFN-g Activity": {"IFN-g": 1},
    "Granzyme B Activity": {"Granzyme B": 1},
    "TNF-a Increase": {"TNF-a": 1},
    "TNF-a Decrease": {"TNF-a": -1},
    "IL-10 Increase": {"IL-10": 1},
}

FLOW_SIGNATURES = {
    "Cytotoxic T-cell Activity": {
        "CTL| Freq. of CD45+ (%)": 1,
        "CD8+GranzymeB+ | Freq. of Parent (%)": 1,
        "CTL freq parent": 1,
    },
    "NK-cell mediated Cytotoxicity": {
        "NK+GranzymeB+ | Freq. of Parent (%)": 1
    },
    "NKT-cell  mediated Cytotoxicity": {
        "NKT+GranzymeB+ | Freq. of Parent (%)": 1
    },
    "T-cell Proliferation": {
        "CD8+Ki67+ | Freq. of Parent (%)": 1,
        "CD4+Ki67+ | Freq. of Parent (%)": 1,
    },
    "NK-cell Proliferation": {"NK+Ki67+ | Freq. of Parent (%)": 1},
    "NKT-cell Proliferation": {"NKT+Ki67+ | Freq. of Parent (%)": 1},
    "Regulatory T-cell Activity": {
        "Treg| Freq. of CD45+ (%)": 1,
        "Treg freq parent": 1,
        "CD4+FoxP3+CTLA4+ | Freq. of Parent (%)": 1,
        "CD4+FoxP3+Ki67+ | Freq. of Parent (%)": 1,
    },
    "Macrophage/M2 Polarization": {
        "Macrophage| Freq. of CD45+ (%)": 1,
        "M2 | Freq. of CD45+ (%)": 1,
    },
}

DEFAULT_SHEET_NAMES = {
    "Histo": "",
    "Cytokine": "",
    "Flow": "",
    "NanoString": "",
}

# ============================================================
# PIPELINE FUNCTIONS (ported from the notebook)
# ============================================================

def calculate_ges_scores(assay4: pd.DataFrame, ges_df: pd.DataFrame) -> pd.DataFrame:
    """Mean log2 expression across each Gene Expression Signature's genes."""
    ges_scores = assay4.iloc[:, :2].copy()  # Sample_ID, Arms

    for ges_name in ges_df.columns:
        genes = ges_df[ges_name].dropna().astype(str).str.strip().tolist()
        genes_present = [g for g in genes if g in assay4.columns]

        if len(genes_present) == 0:
            ges_scores[ges_name] = np.nan
        else:
            expression = assay4[genes_present]
            log2_expression = np.log2(expression)
            ges_scores[ges_name] = log2_expression.mean(axis=1, skipna=True)

    ges_scores.iloc[:, 2:] = ges_scores.iloc[:, 2:].round(3)
    return ges_scores


def calculate_log2fc_nanostring(
    df: pd.DataFrame, control_arm: str, treatment_arm: str
) -> pd.DataFrame:
    """GES values are already log2, so fold change is a simple subtraction.
    Fold change is always Treatment vs Control (Treatment - Control)."""
    feature_cols = [c for c in df.columns if c not in ["Sample_ID", "Arms"]]

    control = df[df["Arms"] == control_arm].set_index("Sample_ID")[feature_cols]
    treatment = df[df["Arms"] == treatment_arm].set_index("Sample_ID")[feature_cols]

    common_samples = control.index.intersection(treatment.index)
    control = control.loc[common_samples]
    treatment = treatment.loc[common_samples]

    return treatment - control


def calculate_log2fc(df: pd.DataFrame, control_arm: str, treatment_arm: str) -> pd.DataFrame:
    """log2 fold change of Treatment vs Control: log2((Treatment+1)/(Control+1))."""
    feature_cols = [c for c in df.columns if c not in ["Sample_ID", "Arms"]]

    control = df[df["Arms"] == control_arm].set_index("Sample_ID")[feature_cols]
    treatment = df[df["Arms"] == treatment_arm].set_index("Sample_ID")[feature_cols]

    common_samples = control.index.intersection(treatment.index)
    control = control.loc[common_samples]
    treatment = treatment.loc[common_samples]

    return np.log2((treatment + 1) / (control + 1))


def select_responsive_features(log2fc: pd.DataFrame, cutoff: float, min_samples: int = 4):
    response = log2fc.abs() >= cutoff
    responding_count = response.sum(axis=0)
    percentage = responding_count / len(log2fc) * 100
    selected = responding_count >= min_samples

    summary = pd.DataFrame(
        {
            "Responding_samples": responding_count,
            "Percentage": percentage,
            "Selected": selected,
        }
    )

    selected_features = summary.index[summary["Selected"]].tolist()
    return selected_features, summary


def get_direction(log2fc: pd.DataFrame, cutoff: float) -> pd.DataFrame:
    return pd.DataFrame(
        np.where(
            log2fc >= cutoff,
            "Increased",
            np.where(log2fc <= -cutoff, "Decreased", "No meaningful change"),
        ),
        index=log2fc.index,
        columns=log2fc.columns,
    )


def get_active_ges_labels(direction_df: pd.DataFrame) -> pd.Series:
    labels = pd.Series(index=direction_df.index, dtype="object")
    for sample in direction_df.index:
        active_ges = [
            ges
            for ges in direction_df.columns
            if direction_df.loc[sample, ges] == "Increased"
        ]
        labels.loc[sample] = "; ".join(active_ges) if active_ges else "No clear activity"
    return labels


def calculate_activity(direction_df, signatures, cutoff=0.50, require_coverage=False):
    activity_score = pd.DataFrame(
        index=direction_df.index, columns=signatures.keys(), dtype=float
    )
    activity_label = pd.DataFrame(index=direction_df.index, columns=signatures.keys())

    for signature, features in signatures.items():
        if require_coverage:
            minimum_required = max(1, int(np.ceil(0.50 * len(features))))
        else:
            minimum_required = 1

        for sample in direction_df.index:
            supporting = 0
            available = 0

            for feature, expected_direction in features.items():
                if feature not in direction_df.columns:
                    continue

                observed = direction_df.loc[sample, feature]
                if observed == "Increased":
                    observed_direction = 1
                elif observed == "Decreased":
                    observed_direction = -1
                else:
                    observed_direction = 0

                if observed_direction == 0:
                    continue

                available += 1
                if observed_direction == expected_direction:
                    supporting += 1

            if available < minimum_required:
                score = np.nan
                label = "Insufficient data"
            else:
                score = supporting / available
                label = "Active" if score >= cutoff else "No clear activity"

            activity_score.loc[sample, signature] = score
            activity_label.loc[sample, signature] = label

    return activity_score, activity_label


def get_active_labels(activity_label_df: pd.DataFrame) -> pd.Series:
    """Collapse an Active/No clear activity/Insufficient data table into
    one semicolon-joined label per sample, listing only 'Active' signatures."""
    labels = pd.Series(index=activity_label_df.index, dtype="object")
    for sample in activity_label_df.index:
        active = [
            sig
            for sig in activity_label_df.columns
            if activity_label_df.loc[sample, sig] == "Active"
        ]
        labels.loc[sample] = "; ".join(active) if active else "No clear activity"
    return labels


def build_activity_binary(final_activity: pd.DataFrame) -> pd.DataFrame:
    """1 = signature present for that sample (in any assay), 0 = absent."""
    assay_columns = list(final_activity.columns)

    unique_labels = set()
    for col in assay_columns:
        for value in final_activity[col].dropna():
            if value != "No clear activity":
                unique_labels.update(x.strip() for x in value.split(";"))

    activity_binary = pd.DataFrame(index=final_activity.index)

    for label in sorted(unique_labels):
        activity_binary[label] = 0
        for sample in final_activity.index:
            for col in assay_columns:
                value = final_activity.loc[sample, col]
                if pd.notna(value):
                    labels = [x.strip() for x in value.split(";")]
                    if label in labels:
                        activity_binary.loc[sample, label] = 1
                        break

    activity_binary.index.name = "Sample_ID"
    return activity_binary


def _elementwise(df: pd.DataFrame, func):
    """DataFrame.applymap was removed in pandas 3.0 in favor of .map; this
    works across both old and new pandas versions."""
    if hasattr(df, "map"):
        try:
            return df.map(func)
        except TypeError:
            pass
    return df.applymap(func)


def format_binary_for_display(activity_binary: pd.DataFrame) -> pd.DataFrame:
    """Presentation copy: 1 stays as 1 ('present'), 0 becomes a blank cell."""
    return _elementwise(activity_binary, lambda v: 1 if v == 1 else "")


def get_available_arms(excel_file: pd.ExcelFile, sheet_name: str) -> list:
    """Read just the 'Arms' column of a sheet to list the distinct study arms."""
    try:
        arms = excel_file.parse(sheet_name, usecols=["Arms"])["Arms"]
    except Exception:
        arms = excel_file.parse(sheet_name)["Arms"]
    return sorted(arms.dropna().astype(str).unique().tolist())


def to_excel_bytes(df: pd.DataFrame) -> bytes:
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as writer:
        df.to_excel(writer)
    return buf.getvalue()


# ------------------------------------------------------------
# Categorized (Y) output layout, modelled on the reporting template
# (H = histopathology, F = flow cytometry, C = cytokine, GES = NanoString)
# Each entry: (category, [(display column, [internal signature names])])
# ------------------------------------------------------------
TEMPLATE_LAYOUT = [
    ("Tumor cytotoxicity", [
        ("Tumor content decrease/caspase 3 increase/Necrosis", ["Tumor cytotoxicity"]),
    ]),
    ("Immune hot TME", [
        ("TIS (GES)", ["TIS"]),
        ("Immune infiltration (H)", ["Immune infiltration"]),
    ]),
    ("Immunosuppressive TME", [
        ("MDSC (GES)", ["MDSC"]),
        ("IL1beta (GES)", ["IL1beta"]),
        ("Matrix_Remodeling_and_Metastasis (GES)", ["Matrix_Remodeling_and_Metastasis"]),
    ]),
    ("T cell activity", [
        ("T-cell Proliferation (F)", ["T-cell Proliferation"]),
        ("Lymphoid_Compartment (GES)", ["Lymphoid_Compartment"]),
    ]),
    ("Cytotoxic T cell activity", [
        ("CD8_Cytotoxic_T_cells (GES)", ["CD8_Cytotoxic_T_cells"]),
        ("Cytotoxic T-cell Activity (F)", ["Cytotoxic T-cell Activity"]),
        ("IFN-g release (C)", ["IFN-g Activity"]),
        ("Granzyme B release (C)", ["Granzyme B Activity"]),
    ]),
    ("T regulatory cell (Treg) activity", [
        ("Treg (GES)", ["Treg"]),
        ("Treg proportion and proliferation (F)", ["Regulatory T-cell Activity"]),
        ("IL-10 increase (C)", ["IL-10 Increase"]),
    ]),
    ("Macrophage activity", [
        ("M1 (GES)", ["M1"]),
        ("M2 (GES)", ["M2"]),
        ("Macrophages (F)", ["Macrophage/M2 Polarization"]),
    ]),
    ("NK cell activity", [
        ("NK_cytotoxic_cells (GES)", ["NK_cytotoxic_cells"]),
    ]),
    ("T helper cell activity", [
        ("Th1 (GES)", ["Th1"]),
        ("Th2 (GES)", ["Th2"]),
    ]),
    ("B-cell activity", [
        ("B-cell (GES)", ["B-cell"]),
    ]),
]

OTHER_GROUP = "Other signatures (not in template)"
PLATFORM_COL = "Platform Response"


def _norm(text: str) -> str:
    """Normalize a signature name for matching (case/space/punctuation-insensitive)."""
    return re.sub(r"[^a-z0-9]", "", str(text).lower())


def build_categorized_output(activity_binary: pd.DataFrame) -> pd.DataFrame:
    """Reshape the flat 1/0 table into the categorized template: 'Y' where the
    signature is present (1), blank otherwise. Signatures that are not part of
    the template are kept under an 'Other' group so nothing is dropped."""
    norm_to_col = {_norm(c): c for c in activity_binary.columns}
    used = set()

    data = {(PLATFORM_COL, ""): pd.Series("", index=activity_binary.index)}

    for category, columns in TEMPLATE_LAYOUT:
        for display_name, keys in columns:
            present = pd.Series(False, index=activity_binary.index)
            for key in keys:
                col = norm_to_col.get(_norm(key))
                if col is not None:
                    used.add(col)
                    present = present | (activity_binary[col] == 1)
            data[(category, display_name)] = present.map(lambda v: "Y" if v else "")

    for col in activity_binary.columns:
        if col not in used:
            data[(OTHER_GROUP, col)] = (activity_binary[col] == 1).map(
                lambda v: "Y" if v else ""
            )

    out = pd.DataFrame(data)
    out.columns = pd.MultiIndex.from_tuples(out.columns)
    out.index.name = "Sample_ID"
    return out


def safe_filename(name: str, default: str = "Activity_binary_NS") -> str:
    """Make a user-typed name safe for use as a file name."""
    cleaned = re.sub(r"[^A-Za-z0-9._-]+", "_", str(name).strip()).strip("._")
    return cleaned or default


def stamped_filename(base: str, stamp: str, ext: str = "xlsx") -> str:
    return f"{safe_filename(base)}_{stamp}.{ext}"


def categorized_to_excel_bytes(cat_df: pd.DataFrame, sheet_name: str) -> bytes:
    """Write the categorized table with a merged two-level header (category
    over signature), styled like the reporting template using Farcast Brand colors."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = re.sub(r"[\[\]:*?/\\]", "_", sheet_name)[:31] or "Activity"

    thin = Side(style="thin", color="CBD5E1")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cat_font = Font(bold=True, color="FFFFFF")
    head_font = Font(bold=True, color="1E2859")
    cat_fill = PatternFill("solid", fgColor="1E2859")  # Farcast Navy
    head_fill = PatternFill("solid", fgColor="F3EDF8")  # Soft Farcast Purple tint
    y_fill = PatternFill("solid", fgColor="F7EBFB")     # Farcast highlight
    y_font = Font(bold=True, color="7C1F92")

    columns = list(cat_df.columns)

    # --- Header: Sample_ID (merged over 2 rows) ---
    ws.cell(row=1, column=1, value="Sample_ID")
    ws.merge_cells(start_row=1, start_column=1, end_row=2, end_column=1)

    # --- Header: category row (merged spans) + signature row ---
    col_idx = 2
    i = 0
    while i < len(columns):
        category = columns[i][0]
        j = i
        while j + 1 < len(columns) and columns[j + 1][0] == category:
            j += 1
        start, end = col_idx + i, col_idx + j
        if columns[i][1] == "":  # single stand-alone header (e.g. Platform Response)
            ws.cell(row=1, column=start, value=category)
            ws.merge_cells(start_row=1, start_column=start, end_row=2, end_column=start)
        else:
            ws.cell(row=1, column=start, value=category)
            if end > start:
                ws.merge_cells(start_row=1, start_column=start, end_row=1, end_column=end)
            for k in range(i, j + 1):
                ws.cell(row=2, column=col_idx + k, value=columns[k][1])
        i = j + 1

    # --- Data rows ---
    for r, (sample_id, row) in enumerate(cat_df.iterrows(), start=3):
        ws.cell(row=r, column=1, value=sample_id)
        for c, value in enumerate(row.tolist(), start=2):
            ws.cell(row=r, column=c, value=(value if value != "" else None))

    n_rows = len(cat_df) + 2
    n_cols = len(columns) + 1

    # --- Styling ---
    for r in range(1, n_rows + 1):
        for c in range(1, n_cols + 1):
            cell = ws.cell(row=r, column=c)
            cell.border = border
            cell.alignment = center
            if r == 1:
                cell.font, cell.fill = cat_font, cat_fill
            elif r == 2:
                cell.font, cell.fill = head_font, head_fill
            elif c == 1:
                cell.font = Font(bold=True, color="1E2859")
            elif cell.value == "Y":
                cell.font, cell.fill = y_font, y_fill

    ws.row_dimensions[1].height = 32
    ws.row_dimensions[2].height = 62
    ws.column_dimensions["A"].width = 16
    for c in range(2, n_cols + 1):
        ws.column_dimensions[get_column_letter(c)].width = 17
    ws.freeze_panes = "C3"

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ============================================================
# STREAMLIT UI
# ============================================================

st.markdown(
    """
    <div class="farcast-header-card">
        <div class="farcast-header-brand">🧬 FARCAST BIOSCIENCES &bull; CANCER IMMUNOTHERAPY PLATFORM</div>
        <h1 class="farcast-header-title">Sample Activity Labeler</h1>
        <p class="farcast-header-desc">
            Multimodal assay integration &bull; Directional response calling &bull; Activity presence matrix
        </p>
        <div class="farcast-badges">
            <span class="farcast-badge-chip">🔬 Histopathology (H&E / IHC)</span>
            <span class="farcast-badge-chip">🧪 Cytokines</span>
            <span class="farcast-badge-chip">📊 Flow Cytometry</span>
            <span class="farcast-badge-chip">🧬 NanoString GES</span>
        </div>
    </div>
    """,
    unsafe_allow_html=True,
)
st.markdown(
    '<p class="app-subtitle">Upload your multimodal assay workbook and gene '
    "signature list in the sidebar, review the analysis parameters, and generate "
    "the publication-ready sample activity presence matrix.</p>",
    unsafe_allow_html=True,
)

with st.sidebar:
    if LOGO_B64:
        st.markdown(
            f"""
            <div class="farcast-sidebar-logo-container">
                <img class="farcast-sidebar-logo" src="data:image/jpeg;base64,{LOGO_B64}" alt="Farcast Biosciences Logo" />
            </div>
            """,
            unsafe_allow_html=True,
        )
    else:
        st.markdown("### **FARCAST** Biosciences")

    st.header("1. Upload data")
    assay_file = st.file_uploader(
        "Multimodal assay Excel workbook (.xlsx)", type=["xlsx"]
    )
    ges_file = st.file_uploader("Gene Signatures list (.csv)", type=["csv"])

    excel_file = None
    available_sheets = []
    if assay_file is not None:
        try:
            excel_file = pd.ExcelFile(assay_file)
            available_sheets = excel_file.sheet_names
        except Exception as e:
            st.error(f"Could not read the workbook: {e}")

    st.header("2. Sheet names")
    st.caption("Confirm which sheet in the workbook holds each assay.")

    def _sheet_picker(label, default_name=""):
        if available_sheets:
            options = [""] + available_sheets
            default_idx = (
                options.index(default_name)
                if default_name in options
                else 0
            )
            return st.selectbox(
                label,
                options,
                index=default_idx,
                format_func=lambda s: "-- Select sheet --" if s == "" else s,
            )
        return st.text_input(label, value=default_name, placeholder="e.g. Enter sheet name")

    sheet_histo = _sheet_picker("Histo sheet", DEFAULT_SHEET_NAMES["Histo"])
    sheet_cytokine = _sheet_picker("Cytokine sheet", DEFAULT_SHEET_NAMES["Cytokine"])
    sheet_flow = _sheet_picker("Flow sheet", DEFAULT_SHEET_NAMES["Flow"])
    sheet_nanostring = _sheet_picker("NanoString sheet", DEFAULT_SHEET_NAMES["NanoString"])

    st.header("3. Arm selection")
    st.caption(
        "Fold change is always calculated as Treatment vs. Control "
        "(log2(Treatment/Control))."
    )

    available_arms = []
    if excel_file is not None and sheet_histo:
        try:
            available_arms = get_available_arms(excel_file, sheet_histo)
        except Exception as e:
            st.warning(f"Could not read 'Arms' column from '{sheet_histo}': {e}")

    if available_arms:
        control_default = 0
        treatment_default = 1 if len(available_arms) > 1 else 0
        # Prefer sensible defaults if the classic Rx-A / Rx-B labels exist
        if "Rx-A" in available_arms:
            control_default = available_arms.index("Rx-A")
        if "Rx-B" in available_arms:
            treatment_default = available_arms.index("Rx-B")

        control_arm = st.selectbox(
            "Control arm (baseline)", available_arms, index=control_default
        )
        treatment_arm = st.selectbox(
            "Treatment arm", available_arms, index=treatment_default
        )
        if control_arm == treatment_arm:
            st.error("Control and Treatment arms must be different.")
    else:
        st.caption("Upload the workbook and select sheets to detect the available arms automatically.")
        control_arm = st.text_input("Control arm (baseline)", "Rx-A")
        treatment_arm = st.text_input("Treatment arm", "Rx-B")

    st.header("4. Parameters")
    fc_cutoff = st.number_input("Fold-change cutoff", value=1.2, step=0.1, min_value=1.0)
    min_samples = st.number_input(
        "Min. responding samples for feature selection", value=4, step=1, min_value=1
    )
    activity_cutoff_histo = st.slider("Histo activity cutoff", 0.0, 1.0, 0.50, 0.05)
    activity_cutoff_cytokine = st.slider(
        "Cytokine activity cutoff", 0.0, 1.0, 1.00, 0.05
    )
    activity_cutoff_flow = st.slider("Flow activity cutoff", 0.0, 1.0, 0.50, 0.05)
    require_coverage = st.checkbox(
        "Require >=50% feature coverage per signature", value=False
    )

    st.header("5. Output")
    output_name = st.text_input(
        "Output file name",
        value="Activity_binary_NS",
        help="Used as the download file name and the Excel sheet name. "
        "The date and time are added automatically.",
    )
    st.caption(
        "Saved as: "
        f"`{stamped_filename(output_name, datetime.now().strftime('%Y%m%d_%H%M%S'))}`"
    )

    sheets_selected = bool(sheet_histo and sheet_cytokine and sheet_flow and sheet_nanostring)
    arms_valid = bool(control_arm and treatment_arm and control_arm != treatment_arm)
    can_run = sheets_selected and arms_valid
    run_button = st.button(
        "Run pipeline",
        type="primary",
        use_container_width=True,
        disabled=not can_run,
    )

if not assay_file or not ges_file:
    st.info("Upload both the assay workbook and the gene signature CSV in the sidebar to begin.")
    st.stop()

if not sheets_selected:
    st.info("Please select the sheet names for all 4 assays (Histo, Cytokine, Flow, NanoString) in the sidebar to proceed.")
    st.stop()

if not run_button and "final_activity" not in st.session_state:
    st.info(
        f"Control = **{control_arm}**, Treatment = **{treatment_arm}**. "
        "Adjust arms/parameters in the sidebar, then click **Run pipeline**."
    )
    st.stop()

if run_button:
    with st.spinner("Running pipeline..."):
        log2fc_cutoff = np.log2(fc_cutoff)

        # 1. Load assays (reuse the already-opened workbook so the uploaded
        # file's stream position doesn't get consumed/confused)
        workbook = excel_file if excel_file is not None else pd.ExcelFile(assay_file)
        assay1 = workbook.parse(sheet_histo)
        assay2 = workbook.parse(sheet_cytokine)
        assay3 = workbook.parse(sheet_flow)
        assay4 = workbook.parse(sheet_nanostring)

        ges_df = pd.read_csv(ges_file)

        # 2. GES scores
        ges_scores = calculate_ges_scores(assay4, ges_df)

        # 3. Fold change (Treatment vs. Control, as selected in the sidebar)
        histo_fc = calculate_log2fc(assay1, control_arm, treatment_arm)
        cytokine_fc = calculate_log2fc(assay2, control_arm, treatment_arm)
        flow_fc = calculate_log2fc(assay3, control_arm, treatment_arm)
        nanostring_fc = calculate_log2fc_nanostring(ges_scores, control_arm, treatment_arm)

        # 4. Feature selection
        selected1, summary1 = select_responsive_features(histo_fc, log2fc_cutoff, min_samples)
        selected2, summary2 = select_responsive_features(cytokine_fc, log2fc_cutoff, min_samples)
        selected3, summary3 = select_responsive_features(flow_fc, log2fc_cutoff, min_samples)
        selected4, summary4 = select_responsive_features(nanostring_fc, log2fc_cutoff, min_samples)

        sel1 = histo_fc[selected1].copy()
        sel2 = cytokine_fc[selected2].copy()
        sel3 = flow_fc[selected3].copy()
        sel4 = nanostring_fc[selected4].copy()

        combined = pd.concat([sel1, sel2, sel3, sel4], axis=1, join="inner")
        combined.index.name = "Sample_ID"

        # 5. Direction
        direction1 = get_direction(sel1, log2fc_cutoff)
        direction2 = get_direction(sel2, log2fc_cutoff)
        direction3 = get_direction(sel3, log2fc_cutoff)
        direction4 = get_direction(sel4, log2fc_cutoff)

        # 6. Signature activity per assay
        histo_score, histo_activity = calculate_activity(
            direction1, HISTO_SIGNATURES, cutoff=activity_cutoff_histo,
            require_coverage=require_coverage,
        )
        cytokine_score, cytokine_activity = calculate_activity(
            direction2, CYTOKINE_SIGNATURES, cutoff=activity_cutoff_cytokine,
            require_coverage=require_coverage,
        )
        flow_score, flow_activity = calculate_activity(
            direction3, FLOW_SIGNATURES, cutoff=activity_cutoff_flow,
            require_coverage=require_coverage,
        )
        nanostring_activity = get_active_ges_labels(direction4)

        # 7. Final activity per assay (one label per sample)
        final_activity = pd.DataFrame(
            {
                "Histo": get_active_labels(histo_activity),
                "Cytokine": get_active_labels(cytokine_activity),
                "Flow": get_active_labels(flow_activity),
                "NanoString": nanostring_activity,
            }
        )
        final_activity.index.name = "Sample_ID"

        # 8. Binary presence table (1 = present)
        activity_binary = build_activity_binary(final_activity)

        # stash everything in session state so results survive reruns
        st.session_state.update(
            dict(
                ges_scores=ges_scores,
                nanostring_fc=nanostring_fc,
                combined=combined,
                summary1=summary1, summary2=summary2, summary3=summary3, summary4=summary4,
                final_activity=final_activity,
                activity_binary=activity_binary,
                control_arm=control_arm,
                treatment_arm=treatment_arm,
                run_stamp=datetime.now().strftime("%Y%m%d_%H%M%S"),
            )
        )
    st.success("Pipeline complete.", icon="✅")

# ============================================================
# RESULTS
# ============================================================

final_activity = st.session_state["final_activity"]
activity_binary = st.session_state["activity_binary"]

st.divider()
st.header("Final output — Activity presence table")
st.caption(
    f"Comparison: **{st.session_state.get('treatment_arm', treatment_arm)}** "
    f"(Treatment) vs. **{st.session_state.get('control_arm', control_arm)}** (Control). "
    "One row per sample, grouped by biological activity. "
    "**Y** means that signature was called Active in at least one assay "
    "for that sample; cells are left blank where it was not detected."
)

n_samples = activity_binary.shape[0]
n_signatures = activity_binary.shape[1]

k1, k2 = st.columns(2)
k1.metric("Samples", n_samples)
k2.metric("Signatures detected", n_signatures)

XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
run_stamp = st.session_state.get("run_stamp", datetime.now().strftime("%Y%m%d_%H%M%S"))
base_name = safe_filename(output_name)


def _stamped(base: str) -> str:
    return stamped_filename(base, run_stamp)


def _style_cells(df: pd.DataFrame, fn):
    styler = df.style
    return styler.map(fn) if hasattr(styler, "map") else styler.applymap(fn)


# Categorized table: Y = present, blank = not present
categorized = build_categorized_output(activity_binary)
_y_style = lambda v: (
    "background-color: #F7EBFB; color: #7C1F92; font-weight: 700; text-align: center;"
    if v == "Y"
    else ""
)
st.dataframe(_style_cells(categorized, _y_style), use_container_width=True)

st.download_button(
    f"Download {_stamped(base_name)}",
    data=categorized_to_excel_bytes(categorized, sheet_name=base_name),
    file_name=_stamped(base_name),
    mime=XLSX_MIME,
    use_container_width=True,
    type="primary",
)

with st.expander("Flat binary matrix (1 = present, blank = absent)"):
    display_binary = format_binary_for_display(activity_binary)
    _flat_style = lambda v: (
        "background-color: #F7EBFB; color: #7C1F92; font-weight: 700; text-align: center;"
        if v == 1
        else ""
    )
    st.dataframe(_style_cells(display_binary, _flat_style), use_container_width=True)
    st.download_button(
        f"Download {_stamped(base_name + '_flat')}",
        data=to_excel_bytes(display_binary),
        file_name=_stamped(base_name + "_flat"),
        mime=XLSX_MIME,
        key="dl_flat_binary",
    )

with st.expander("Per-assay activity labels (Final_activity_NS)"):
    st.dataframe(final_activity, use_container_width=True)
    st.download_button(
        f"Download {_stamped('Final_activity_NS')}",
        data=to_excel_bytes(final_activity),
        file_name=_stamped("Final_activity_NS"),
        mime=XLSX_MIME,
        key="dl_final_activity",
    )

with st.expander("Intermediate results"):
    st.subheader("GES scores")
    st.dataframe(st.session_state["ges_scores"], use_container_width=True)
    st.download_button(
        f"Download {_stamped('GES_scores')}",
        data=to_excel_bytes(st.session_state["ges_scores"]),
        file_name=_stamped("GES_scores"),
        key="dl_ges",
    )

    st.subheader("NanoString log2FC")
    st.dataframe(st.session_state["nanostring_fc"], use_container_width=True)
    st.download_button(
        f"Download {_stamped('Nanostring_FC')}",
        data=to_excel_bytes(st.session_state["nanostring_fc"]),
        file_name=_stamped("Nanostring_FC"),
        key="dl_ns_fc",
    )

    st.subheader("Combined selected-feature fold changes")
    st.dataframe(st.session_state["combined"], use_container_width=True)
    st.download_button(
        f"Download {_stamped('combined_FC_latest')}",
        data=to_excel_bytes(st.session_state["combined"]),
        file_name=_stamped("combined_FC_latest"),
        key="dl_combined",
    )

    st.subheader("Feature selection summaries")
    tab1, tab2, tab3, tab4 = st.tabs(["Histo", "Cytokine", "Flow", "NanoString"])
    with tab1:
        st.dataframe(st.session_state["summary1"], use_container_width=True)
    with tab2:
        st.dataframe(st.session_state["summary2"], use_container_width=True)
    with tab3:
        st.dataframe(st.session_state["summary3"], use_container_width=True)
    with tab4:
        st.dataframe(st.session_state["summary4"], use_container_width=True)
